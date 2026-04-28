import os

from docx import Document


def debug_positions(file_path):
    print(f"\n--- Debugging: {os.path.basename(file_path)} ---")
    try:
        doc = Document(file_path)
        print(f"Total Paragraphs: {len(doc.paragraphs)}")
        for index, paragraph in enumerate(doc.paragraphs[:50]):
            text = paragraph.text
            print(f"Para {index}: Len {len(text)} | Content: {repr(text[:30])}")
    except Exception as error:
        print(f"Error: {error}")


if __name__ == "__main__":
    if os.path.exists("before.docx"):
        debug_positions("before.docx")
