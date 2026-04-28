import os

from docx import Document


def analyze_docx(file_path):
    print(f"\n--- Analyzing: {os.path.basename(file_path)} ---")
    try:
        doc = Document(file_path)
    except Exception as error:
        print(f"Error opening {file_path}: {error}")
        return

    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    for index, paragraph in enumerate(doc.paragraphs[:100]):
        text = paragraph.text
        if not text.strip():
            if len(text) > 0:
                print(f"  Para {index}: Empty but length {len(text)} -> {repr(text)}")
        else:
            non_printable = "".join(
                character for character in text
                if not character.isprintable() and character not in ("\t", "\n", "\r")
            )
            if non_printable:
                print(f"  Para {index}: Non-printable chars found: {repr(non_printable)}")

    print(f"Total Tables: {len(doc.tables)}")
    for table_index, table in enumerate(doc.tables):
        print(f"  Table {table_index}: {len(table.rows)} rows x {len(table.columns)} columns")
        for row_index, row in enumerate(table.rows[:3]):
            for column_index, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                if not cell_text and len(cell.paragraphs) > 1:
                    print(
                        f"    Table {table_index} Row {row_index} Col {column_index}: "
                        f"{len(cell.paragraphs)} empty paragraphs"
                    )


if __name__ == "__main__":
    analyze_docx("before.docx")
    analyze_docx("after.docx")
